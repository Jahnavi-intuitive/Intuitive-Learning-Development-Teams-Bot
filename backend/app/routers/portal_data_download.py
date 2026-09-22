# TODO(access-control): This tab has NO role-based access control yet. It
# exports row-level, per-employee data (name/email/practice/etc.) — a
# materially more sensitive level than the Dashboard/Trends tabs, which only
# ever show aggregate counts. Fine for now (internal presentation/testing
# group only), but this MUST be gated (e.g. L&D admins / Practice Directors
# only) before it's exposed to anyone beyond that group. See
# DEVELOPMENT_LOG.md's 2026-09-17 entry for the same flag — don't let this
# slip through unaddressed.
"""Data Download tab (frontend/portal.html) — server-generated .xlsx export
of real row-level certification data. Reads through the exact same
lnd_portal database/queries the Dashboard and Trends tabs use (no separate
calculation path, no duplicated aggregation layer), so an exported number
can never drift from what those tabs show — this endpoint only adds
row-level detail and column selection on top of the same source of truth.

Two modes, confirmed with the user:
- provider: one workbook, one sheet per certification under that provider.
- certification: one workbook, one sheet, just that certification's rows.

Both unfiltered by status (confirmed) and built fully in memory (confirmed
— real data is ~730 rows total across all providers, nowhere near needing
streaming). Filename date is the data-as-of date (confirmed) — the latest
last_synced_at among the actual exported rows, not the day it was clicked.
"""

import io
import re
from datetime import date, datetime, timezone

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.portal_db import get_portal_db
from app.portal_models import CertificationRecord, Provider, StatusCanonical

router = APIRouter(prefix="/api/portal/data-download", tags=["portal-data-download"])

# Server-side display labels, mirroring frontend/portal.html's PROVIDER_META —
# kept in sync by hand (small, stable list); used for filenames and the
# provider dropdown label so "gemini" reads as "Google" here too.
_PROVIDER_LABELS = {
    Provider.aws: "AWS", Provider.anthropic: "Anthropic", Provider.azure: "Azure",
    Provider.databricks: "Databricks", Provider.gemini: "Google", Provider.snowflake: "Snowflake",
}

_STATUS_LABELS = {
    StatusCanonical.completed: "Completed",
    StatusCanonical.in_progress: "In Progress",
    StatusCanonical.not_cleared: "Attempted - Not Cleared",
    StatusCanonical.unknown: "Unknown",
}

# Fixed field order (confirmed with the user) — export columns always follow
# this order regardless of the order fields were checked/unchecked in the UI.
_FIELD_DEFS: list[tuple[str, str, callable]] = [
    ("employee_name", "Employee Name", lambda r: r.person_name),
    ("email", "Email", lambda r: r.person_email),
    ("practice", "Practice", lambda r: r.practice_canonical),
    ("manager", "Manager", lambda r: r.manager_name),
    ("practice_director", "Practice Director", lambda r: r.pd_leader_name),
    ("certification_name", "Certification Name", lambda r: r.certification_name),
    ("provider", "Provider", lambda r: _PROVIDER_LABELS.get(r.provider, r.provider.value)),
    ("status", "Status", lambda r: _STATUS_LABELS.get(r.status_canonical, r.status_canonical.value)),
    ("prep_start_date", "Prep Start Date", lambda r: r.prep_start_date),
    ("exam_date", "Exam Date", lambda r: r.exam_date),
    ("submission_date", "Submission Date", lambda r: r.submission_date),
    ("completion_date", "Completion Date", lambda r: r.completion_date),
    ("target_date", "Target Date", lambda r: r.target_date),
    ("expiry_date", "Expiry Date", lambda r: r.expiry_date),
    ("certificate_number", "Certificate Number", lambda r: r.certificate_number),
    ("comments", "Comments", lambda r: r.comments),
]
_FIELD_KEYS = {key for key, _, _ in _FIELD_DEFS}


class FieldOut(BaseModel):
    key: str
    label: str


class ProviderOptionOut(BaseModel):
    provider: str
    label: str


class CertificationOptionOut(BaseModel):
    provider: str
    provider_label: str
    certification_name: str


@router.get("/fields", response_model=list[FieldOut])
def list_fields():
    return [FieldOut(key=key, label=label) for key, label, _ in _FIELD_DEFS]


@router.get("/providers", response_model=list[ProviderOptionOut])
def list_providers(db: Session = Depends(get_portal_db)):
    """Only providers with at least one real record — nothing to export for
    a provider with zero rows."""
    providers_with_data = db.execute(select(CertificationRecord.provider).distinct()).scalars().all()
    ordered = [p for p in Provider if p in set(providers_with_data)]
    return [ProviderOptionOut(provider=p.value, label=_PROVIDER_LABELS.get(p, p.value)) for p in ordered]


@router.get("/certifications", response_model=list[CertificationOptionOut])
def list_certifications(db: Session = Depends(get_portal_db)):
    """Every distinct (provider, certification_name) with at least one row,
    any status — unfiltered, matching the export itself. Sorted by
    provider's catalog order then certification name, so the frontend's
    grouped dropdown renders in a stable, scannable order."""
    rows = db.execute(
        select(CertificationRecord.provider, CertificationRecord.certification_name)
        .where(CertificationRecord.certification_name.is_not(None))
        .distinct()
    ).all()
    provider_order = {p: i for i, p in enumerate(Provider)}
    rows.sort(key=lambda r: (provider_order[r[0]], r[1]))
    return [
        CertificationOptionOut(provider=p.value, provider_label=_PROVIDER_LABELS.get(p, p.value), certification_name=name)
        for p, name in rows
    ]


_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = _INVALID_SHEET_CHARS.sub("-", name or "Sheet").strip() or "Sheet"
    cleaned = cleaned[:31]
    base, i = cleaned, 2
    while cleaned.lower() in used:
        suffix = f" ({i})"
        cleaned = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(cleaned.lower())
    return cleaned


_FILENAME_INVALID = re.compile(r"[^A-Za-z0-9_-]+")


def _safe_filename_part(name: str) -> str:
    return _FILENAME_INVALID.sub("_", name).strip("_") or "export"


_HEADER_FILL = PatternFill(fill_type="solid", start_color="FF1F4E78", end_color="FF1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFFFF")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
_MIN_COL_WIDTH = 10
_MAX_COL_WIDTH = 45


def _write_sheet(wb: Workbook, sheet_name: str, rows: list[CertificationRecord], fields: list[tuple[str, str, callable]], first: bool) -> None:
    ws = wb.active if first else wb.create_sheet()
    ws.title = sheet_name
    labels = [label for _, label, _ in fields]
    ws.append(labels)
    sorted_rows = sorted(rows, key=lambda r: (r.person_name or "").lower())
    for r in sorted_rows:
        ws.append([_cell_value(extractor(r)) for _, _, extractor in fields])

    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGNMENT
    ws.freeze_panes = "A2"

    # Column width sized to its longest real value (header or data), not a
    # fixed default — the field list mixes short codes (Status) and long
    # free text (Comments), so one width for all of them is always wrong.
    for col_idx, label in enumerate(labels, start=1):
        longest = len(str(label))
        for r in sorted_rows:
            value = _cell_value(fields[col_idx - 1][2](r))
            if value is not None:
                longest = max(longest, len(str(value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(_MIN_COL_WIDTH, min(longest + 3, _MAX_COL_WIDTH))


def _cell_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _data_as_of(rows: list[CertificationRecord]) -> str:
    timestamps = [r.last_synced_at for r in rows if r.last_synced_at is not None]
    as_of = max(timestamps) if timestamps else datetime.now(timezone.utc)
    return as_of.date().isoformat()


def _resolve_fields(field_keys: list[str] | None) -> list[tuple[str, str, callable]]:
    if not field_keys:
        return list(_FIELD_DEFS)
    unknown = set(field_keys) - _FIELD_KEYS
    if unknown:
        raise HTTPException(status_code=422, detail=f"Unknown field(s): {sorted(unknown)}")
    selected = set(field_keys)
    return [f for f in _FIELD_DEFS if f[0] in selected]


@router.get("/export/provider")
def export_provider(
    provider: Provider,
    fields: list[str] | None = Query(None),
    db: Session = Depends(get_portal_db),
):
    """One workbook, one sheet per certification under `provider`, full
    row-level data (every employee record), unfiltered by status."""
    field_defs = _resolve_fields(fields)
    rows = db.execute(select(CertificationRecord).where(CertificationRecord.provider == provider)).scalars().all()
    if not rows:
        raise HTTPException(status_code=404, detail=f"No records for provider {provider.value!r}.")

    by_cert: dict[str, list[CertificationRecord]] = {}
    for r in rows:
        by_cert.setdefault(r.certification_name or "(Unspecified)", []).append(r)

    wb = Workbook()
    used_names: set[str] = set()
    for i, (cert_name, cert_rows) in enumerate(sorted(by_cert.items())):
        _write_sheet(wb, _safe_sheet_name(cert_name, used_names), cert_rows, field_defs, first=(i == 0))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    label = _PROVIDER_LABELS.get(provider, provider.value)
    filename = f"{_safe_filename_part(label)}_certifications_{_data_as_of(rows)}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export/certification")
def export_certification(
    provider: Provider,
    certification_name: str,
    fields: list[str] | None = Query(None),
    db: Session = Depends(get_portal_db),
):
    """One workbook, one sheet, just this certification's row-level data."""
    field_defs = _resolve_fields(fields)
    rows = db.execute(
        select(CertificationRecord).where(
            CertificationRecord.provider == provider, CertificationRecord.certification_name == certification_name
        )
    ).scalars().all()
    if not rows:
        raise HTTPException(status_code=404, detail=f"No records for {provider.value}/{certification_name!r}.")

    wb = Workbook()
    _write_sheet(wb, _safe_sheet_name(certification_name, set()), rows, field_defs, first=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"{_safe_filename_part(certification_name)}_{_data_as_of(rows)}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
